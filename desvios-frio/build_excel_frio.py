import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

import os

STOCK_SRC = '/root/.claude/uploads/63357b75-0e3d-5611-bced-932fcb8f796a/0076f13b-Informe_Stock_Pa_s_20260820.xlsx'
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'Desvios_de_Frio.xlsx')
# Histórico acumulado por build_data_frio.py (mismo criterio de upsert por SKU/Semana/CADENA) —
# se usa como fuente en vez de re-leer el Excel crudo, para que este reporte siempre refleje
# todas las semanas acumuladas y no solo lo que traiga el último archivo subido.
HIST_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'raw_historico_frio.csv')

# ══════════════════════════════════════════════════════════════════
# Carga (ya viene limpio de build_data_frio.py, que debe correrse antes que este script)
# ══════════════════════════════════════════════════════════════════
df = pd.read_csv(HIST_PATH)

semanas = sorted(df['Semana'].unique())


def sem_label(s):
    s = int(s)
    return f"{str(s // 100)[2:]}-S{s % 100:02d}"


# ── Exclusión para las hojas de análisis (no para la Base): FCST=0 en >40% de semanas activas ──
wk_sku_fcst = df.groupby(['SKU', 'Semana'])['FCST'].sum().reset_index()


def _stats(g):
    return pd.Series({'n': len(g), 'nz': (g['FCST'] <= 0.001).sum()})


stats = wk_sku_fcst.groupby('SKU').apply(_stats, include_groups=False)
stats['pct'] = stats['nz'] / stats['n']
sku_excluidos = stats[stats['pct'] > 0.4].index
# Excepción manual: incluir igual estos SKU aunque superen el umbral (pedido puntual).
SKU_FORZAR_INCLUSION = {30002120}
sku_excluidos = sku_excluidos.difference(SKU_FORZAR_INCLUSION)
df_an = df[~df['SKU'].isin(sku_excluidos)].copy()  # universo para hojas de análisis

# ══════════════════════════════════════════════════════════════════
# Agregado por SKU (todo el histórico, todas las cadenas)
# ══════════════════════════════════════════════════════════════════
g = df_an.groupby(['SKU', 'Nombre Producto', 'Marca', 'Categoria Producto', 'SubCat DMD']).agg(
    FCST=('FCST', 'sum'), Solicitado=('Solicitado', 'sum'), SellIn=('Venta Sell IN', 'sum'),
    VentaReal=('Venta Real', 'sum'), Quebrados=('Quebrados', 'sum'),
).reset_index()
g = g.rename(columns={'Categoria Producto': 'Categoria', 'SubCat DMD': 'SubCat'})
g['gap_FS_t'] = g['Solicitado'] - g['FCST']
g['gap_FS_pct'] = (g['Solicitado'] / g['FCST'].replace(0, pd.NA) - 1) * 100
g['gap_SS_t'] = g['Solicitado'] - g['SellIn']
g['gap_SS_pct'] = (g['Solicitado'] / g['SellIn'].replace(0, pd.NA) - 1) * 100
g['des_pct'] = (g['SellIn'] / g['FCST'].replace(0, pd.NA)) * 100
g['gap_FC_t'] = g['FCST'] - g['SellIn']

tot = g[['FCST', 'Solicitado', 'SellIn', 'VentaReal', 'Quebrados']].sum()
gap_FS = float(tot['Solicitado'] - tot['FCST'])
gap_SS = float(tot['Solicitado'] - tot['SellIn'])
gap_FC = float(tot['FCST'] - tot['SellIn'])

# ── Stock en riesgo / liquidación (match por SKU) ──
stock = pd.read_excel(STOCK_SRC, sheet_name='DETALLE WMS')
stock['es_vliq'] = stock['ESTADO'] == 'VLIQ'
stock['en_riesgo'] = stock['es_vliq'] | (stock['PORCENTAJE'] > 26)
risk_all = stock[stock['en_riesgo']].copy()
sku_universe = set(g['SKU'].tolist())
risk = risk_all[risk_all['CODIGO_SAP'].isin(sku_universe)].copy()
risk_by_sku = risk.groupby(['CODIGO_SAP', 'Nombre Producto']).apply(
    lambda x: pd.Series({
        'ton_riesgo': x['KILOS'].sum() / 1000,
        'ton_vliq': x.loc[x['ESTADO'] == 'VLIQ', 'KILOS'].sum() / 1000,
        'n_lotes': x['LOTE'].nunique(),
    }), include_groups=False
).reset_index() if len(risk) else pd.DataFrame(columns=['CODIGO_SAP', 'Nombre Producto', 'ton_riesgo', 'ton_vliq', 'n_lotes'])
cruce = risk_by_sku.merge(
    g[['SKU', 'Marca', 'Categoria', 'FCST', 'SellIn', 'des_pct', 'gap_FC_t', 'gap_SS_t']],
    left_on='CODIGO_SAP', right_on='SKU', how='left'
)

# ══════════════════════════════════════════════════════════════════
# Estilos
# ══════════════════════════════════════════════════════════════════
wb = Workbook()
RED = 'C8001E'
DARK = '1A2733'
GREY = 'F2F2F2'
WHITE = 'FFFFFF'
header_font = Font(bold=True, color=WHITE, size=10)
header_fill = PatternFill('solid', fgColor=RED)
title_font = Font(bold=True, size=14, color=DARK)
sub_font = Font(size=10, italic=True, color='555555')
thin = Side(style='thin', color='D9D9D9')
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border


def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_df(ws, dfx, start_row, start_col=1, pct_cols=(), ton_cols=()):
    ncols = len(dfx.columns)
    for j, col in enumerate(dfx.columns):
        ws.cell(row=start_row, column=start_col + j, value=col)
    style_header(ws, start_row, start_col + ncols - 1)
    for i, (_, row) in enumerate(dfx.iterrows()):
        r = start_row + 1 + i
        for j, col in enumerate(dfx.columns):
            val = row[col]
            if pd.isna(val) if not isinstance(val, str) else False:
                val = None
            cell = ws.cell(row=r, column=start_col + j, value=val)
            cell.border = border
            if col in pct_cols:
                cell.number_format = '0.0"%"'
            elif col in ton_cols:
                cell.number_format = '#,##0.0'
            if i % 2 == 1:
                cell.fill = PatternFill('solid', fgColor=GREY)
    return start_row + 1 + len(dfx)


# ── HOJA 1: Resumen ──────────────────────────────────────────────
ws = wb.active
ws.title = 'Resumen'
ws['B2'] = 'Desvíos de Frío — FCST / Solicitado / Sell In'
ws['B2'].font = title_font
ws['B3'] = (f'Histórico completo {sem_label(semanas[0])} a {sem_label(semanas[-1])} · {len(g)} SKU activos '
            f'· {len(sku_excluidos)} SKU excluidos (FCST=0 en >40% de sus semanas) · 9 categorías Frío')
ws['B3'].font = sub_font

kpis = [
    ('FCST total (ton)', tot['FCST']), ('Solicitado total (ton)', tot['Solicitado']),
    ('Sell In real (ton)', tot['SellIn']), ('Venta Real (ton)', tot['VentaReal']),
    ('Quebrados (ton)', tot['Quebrados']),
]
r = 5
ws.cell(row=r, column=2, value='KPI').font = Font(bold=True, color=WHITE)
ws.cell(row=r, column=3, value='Toneladas').font = Font(bold=True, color=WHITE)
style_header(ws, r, 3)
for i, (k, v) in enumerate(kpis):
    rr = r + 1 + i
    ws.cell(row=rr, column=2, value=k).border = border
    c = ws.cell(row=rr, column=3, value=round(float(v), 1))
    c.number_format = '#,##0.0'
    c.border = border
    if i % 2 == 1:
        for cc in (2, 3):
            ws.cell(row=rr, column=cc).fill = PatternFill('solid', fgColor=GREY)

r2 = r + len(kpis) + 3
ws.cell(row=r2, column=2, value='Indicador').font = Font(bold=True, color=WHITE)
ws.cell(row=r2, column=3, value='Valor').font = Font(bold=True, color=WHITE)
ws.cell(row=r2, column=4, value='Lectura').font = Font(bold=True, color=WHITE)
style_header(ws, r2, 4)
indicadores = [
    ('des = Sell In / FCST', f"{tot['SellIn']/tot['FCST']*100:.1f}%", 'Exactitud real del forecast vs lo que efectivamente se vendió.'),
    ('Gap Solicitado - FCST', f'{gap_FS:,.0f} ton', 'Toneladas pedidas a producción por encima (o bajo) del forecast.'),
    ('Gap Solicitado - Sell In', f'{gap_SS:,.0f} ton', 'Toneladas producidas/solicitadas que no se colocaron en el canal.'),
    ('Gap FCST - Sell In', f'{gap_FC:,.0f} ton', 'Sobre-forecast puro (si el FCST fuera la única causa del excedente).'),
]
for i, (k, v, lect) in enumerate(indicadores):
    rr = r2 + 1 + i
    ws.cell(row=rr, column=2, value=k).border = border
    ws.cell(row=rr, column=3, value=v).border = border
    cell = ws.cell(row=rr, column=4, value=lect)
    cell.border = border
    cell.alignment = Alignment(wrap_text=True, vertical='center')
    if i % 2 == 1:
        for cc in (2, 3, 4):
            ws.cell(row=rr, column=cc).fill = PatternFill('solid', fgColor=GREY)

nota_row = r2 + len(indicadores) + 2
nota = ws.cell(row=nota_row, column=2, value=(
    'Nota de calidad de datos: 195 de 278 SKU tenían dos nombres de Categoría/SubCategoría distintos en el '
    'histórico (reclasificación del maestro, ej. "MARGARINAS" → "MARGARINA REGULAR"). Se usó la clasificación '
    'más reciente de cada SKU para todo su histórico. "YOGHURT" y "YOGURT" se unificaron en "YOGURT".'))
nota.font = Font(italic=True, color='777777', size=9)
nota.alignment = Alignment(wrap_text=True)
ws.merge_cells(start_row=nota_row, start_column=2, end_row=nota_row + 2, end_column=5)

autosize(ws, [3, 26, 16, 55, 12])
ws.sheet_view.showGridLines = False

# ── HOJAS 2-4: SKU completo (FCST vs Solicitado / Solicitado vs SellIn / FCST vs SellIn) ──
def sku_sheet(name, title, sort_col, cols, headers, pct_cols, ton_cols):
    ws_ = wb.create_sheet(name)
    ws_[f'B2'] = title
    ws_['B2'].font = title_font
    ws_['B3'] = f'{len(g)} SKU · todas las cadenas · histórico completo · ordenado por brecha absoluta'
    ws_['B3'].font = sub_font
    t = g.assign(_abs=g[sort_col].abs()).sort_values('_abs', ascending=False)[cols]
    t.columns = headers
    write_df(ws_, t.round(2), 5, start_col=2, pct_cols=pct_cols, ton_cols=ton_cols)
    autosize(ws_, [3, 10, 32, 16, 18, 12, 15, 12, 12])
    ws_.sheet_view.showGridLines = False
    ws_.freeze_panes = 'B6'


sku_sheet('SKU - FCST vs Solicitado', 'Todos los SKU — FCST vs Solicitado', 'gap_FS_t',
          ['SKU', 'Nombre Producto', 'Marca', 'Categoria', 'FCST', 'Solicitado', 'gap_FS_t', 'gap_FS_pct'],
          ['SKU', 'Producto', 'Marca', 'Categoría', 'FCST (ton)', 'Solicitado (ton)', 'Gap (ton)', 'Gap (%)'],
          ['Gap (%)'], ['FCST (ton)', 'Solicitado (ton)', 'Gap (ton)'])
sku_sheet('SKU - Solicitado vs SellIn', 'Todos los SKU — Solicitado vs Sell In', 'gap_SS_t',
          ['SKU', 'Nombre Producto', 'Marca', 'Categoria', 'Solicitado', 'SellIn', 'gap_SS_t', 'gap_SS_pct'],
          ['SKU', 'Producto', 'Marca', 'Categoría', 'Solicitado (ton)', 'Sell In (ton)', 'Gap (ton)', 'Gap (%)'],
          ['Gap (%)'], ['Solicitado (ton)', 'Sell In (ton)', 'Gap (ton)'])
sku_sheet('SKU - FCST vs SellIn', 'Todos los SKU — FCST vs Sell In (exactitud real)', 'gap_FC_t',
          ['SKU', 'Nombre Producto', 'Marca', 'Categoria', 'FCST', 'SellIn', 'gap_FC_t', 'des_pct'],
          ['SKU', 'Producto', 'Marca', 'Categoría', 'FCST (ton)', 'Sell In (ton)', 'Gap (ton)', 'des % (SellIn/FCST)'],
          ['des % (SellIn/FCST)'], ['FCST (ton)', 'Sell In (ton)', 'Gap (ton)'])

# ── HOJA 5: Stock Riesgo-Liquidacion (detalle) ──────────────────────
ws6 = wb.create_sheet('Stock Riesgo-Liquidacion')
ws6['B2'] = 'Stock en riesgo de liquidación — todas las categorías Frío'
ws6['B2'].font = title_font
ws6['B3'] = (f'Snapshot 20-Ago-2026 · Criterio: ESTADO=VLIQ o %vida útil consumida > 26% · '
             f'{len(risk)} lotes · {risk["KILOS"].sum()/1000:.1f} ton totales')
ws6['B3'].font = sub_font
det = risk[['CODIGO_SAP', 'Nombre Producto', 'CATEGORIA', 'BODEGA', 'LOTE', 'ESTADO', 'PORCENTAJE',
            'FECHA_VENCIMIENTO', 'KILOS']].copy()
det['TON'] = (det['KILOS'] / 1000).round(3)
det = det.drop(columns=['KILOS']).sort_values('TON', ascending=False)
det.columns = ['SKU', 'Producto', 'Categoría Stock', 'Bodega', 'Lote', 'Estado', '% Vida Útil Consumida', 'Fecha Vencimiento', 'Ton']
write_df(ws6, det, 5, start_col=2, pct_cols=[], ton_cols=['Ton'])
autosize(ws6, [3, 10, 34, 16, 12, 9, 14, 16, 16, 9])
ws6.sheet_view.showGridLines = False
ws6.freeze_panes = 'B6'

# ── HOJA 6: Cruce Riesgo x Desviacion ────────────────────────────────
ws7 = wb.create_sheet('Cruce Riesgo x Desviacion')
ws7['B2'] = 'Cruce: Stock en riesgo/liquidación vs. desviación FCST'
ws7['B2'].font = title_font
ws7['B3'] = f'{len(cruce)} SKU con stock en riesgo, cruzados contra el histórico completo de desvíos'
ws7['B3'].font = sub_font
tc = cruce[['CODIGO_SAP', 'Nombre Producto', 'Marca', 'Categoria', 'ton_riesgo', 'ton_vliq', 'FCST', 'SellIn', 'des_pct', 'gap_FC_t']].copy()
tc.columns = ['SKU', 'Producto', 'Marca', 'Categoría', 'Ton en riesgo', 'Ton ya VLIQ', 'FCST (t)', 'Sell In (t)', 'des % (SellIn/FCST)', 'Gap FCST-SellIn (t)']
tc = tc.sort_values('Ton en riesgo', ascending=False)
end_row = write_df(ws7, tc.round(2), 5, start_col=2, pct_cols=['des % (SellIn/FCST)'],
                    ton_cols=['Ton en riesgo', 'Ton ya VLIQ', 'FCST (t)', 'Sell In (t)', 'Gap FCST-SellIn (t)'])
autosize(ws7, [3, 10, 32, 16, 18, 13, 12, 12, 12, 16, 18])
ws7.sheet_view.showGridLines = False
ws7.freeze_panes = 'B6'

nota2_row = end_row + 2
con_forecast = tc.loc[tc['des % (SellIn/FCST)'] < 95, 'Ton en riesgo'].sum()
total_riesgo = tc['Ton en riesgo'].sum()
nota2 = ws7.cell(row=nota2_row, column=2, value=(
    f'Umbral de materialidad: un SKU cuenta como "con sobre-forecast" solo si des% < 95% (SellIn quedó al '
    f'menos 5% bajo el FCST) — un gap apenas positivo en una sola semana no basta. Con este criterio, '
    f'{con_forecast:.1f} de {total_riesgo:.1f} ton en riesgo ({con_forecast/total_riesgo*100 if total_riesgo else 0:.0f}%) '
    f'se explican por sobre-forecast material; el resto tiene otra causa (rotación, distribución, etc.).'))
nota2.font = Font(bold=True, color=DARK, size=10)
nota2.alignment = Alignment(wrap_text=True)
ws7.merge_cells(start_row=nota2_row, start_column=2, end_row=nota2_row + 3, end_column=6)

# ── HOJA 7: Base (datos crudos, lista para Tabla Dinámica + Slicers) ──
wsb = wb.create_sheet('Base')
base_cols = ['Mes', 'Semana', 'CADENA', 'SKU', 'Nombre Producto', 'Marca', 'Categoria Producto', 'SubCat DMD',
             'FCST', 'Solicitado', 'Venta Sell IN', 'Venta Real', 'Venta Sell OUT', 'Quebrados', 'Bloqueados']
base = df[base_cols].copy()
base['Mes'] = base['Mes'].astype(str)
for j, col in enumerate(base_cols):
    wsb.cell(row=1, column=1 + j, value=col)
style_header(wsb, 1, len(base_cols))
for i, row in enumerate(base.itertuples(index=False), start=2):
    for j, val in enumerate(row):
        wsb.cell(row=i, column=1 + j, value=val)
last_col_letter = get_column_letter(len(base_cols))
last_row = len(base) + 1
tbl = Table(displayName='BaseDesvios', ref=f'A1:{last_col_letter}{last_row}')
tbl.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
wsb.add_table(tbl)
autosize(wsb, [10, 10, 20, 10, 34, 18, 18, 22, 10, 10, 12, 12, 12, 10, 10])
wsb.freeze_panes = 'A2'
wsb.sheet_view.showGridLines = False

wb.save(OUT)
print('OK ->', OUT)
print('SKUs analisis:', len(g), '| excluidos:', len(sku_excluidos), '| filas Base:', len(base))
print('Stock riesgo SKU match:', len(cruce), 'ton:', cruce['ton_riesgo'].sum())
