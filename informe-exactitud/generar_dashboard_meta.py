"""
Genera dashboard_exactitud.html a partir de "Exactitud_Meta_Mensual.xlsx".

Este archivo reemplaza al metodo anterior (CSV por semana desde tablas
dinamicas). Es una tabla plana, una fila por SKU x Mes x Tipo Indicador:

    Mes | CPFR | Cadena | SKU | Nombre | Categoria Producto |
    Tipo Indicador (MENSUAL/SEMANAL) | Meta | Venta SI | Error Abs | Exactitud

Notas importantes (verificadas con los datos, no asumidas):
- CPFR = Cadena + Tipo de Almacenamiento (ej. "WALMART FRIO", "TOTTUS SECO").
  Canal Tradicional y Supermercados Region no se dividen en Frio/Seco.
- "Tipo Indicador" MENSUAL vs SEMANAL NO son duplicados de la misma cadena:
  cada CPFR usa casi exclusivamente uno de los dos (ej. "ALVI FRIO" viene
  siempre como SEMANAL, "ALVI SECO" siempre como MENSUAL) - es solo la
  cadencia con la que esa cadena/categoria actualiza su meta. Para totales
  de mes hay que SUMAR ambos tipos, nunca filtrar por uno solo, o se
  pierden cadenas enteras.
- Este archivo NO trae detalle por semana individual dentro del mes -
  solo por mes. El ultimo mes disponible esta en curso (parcial).
- Formula de exactitud agregada (confirmada por el usuario, misma que la
  columna Exactitud por fila): Exact = max(0, 1 - sum(Error Abs)/sum(Meta)).
  Desv = (sum(Venta SI) - sum(Meta)) / sum(Meta).

Uso:
    python3 generar_dashboard_meta.py "Exactitud_Meta_Mensual.xlsx"
"""
import json
import sys
import warnings
from pathlib import Path
from collections import defaultdict

warnings.filterwarnings("ignore")
from openpyxl import load_workbook

HERE = Path(__file__).parent
OUT_HTML = HERE / "dashboard_exactitud.html"
META_OBJETIVO = 0.70  # ajustar aqui si cambia la meta corporativa


def leer_filas(path):
    wb = load_workbook(path, data_only=True, read_only=True, keep_links=False)
    ws = wb[wb.sheetnames[0]]
    filas = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        mes = row[0]
        if not mes:
            continue
        filas.append({
            "mes": str(mes), "cpfr": row[1], "cadena": row[2], "sku": row[3],
            "nombre": row[4], "categoria": row[5], "tipo_ind": row[6],
            "meta": row[7] or 0, "si": row[8] or 0, "err": row[9] or 0,
        })
    return filas


def agg(filas):
    meta = sum(f["meta"] for f in filas)
    si = sum(f["si"] for f in filas)
    err = sum(f["err"] for f in filas)
    exact = max(0.0, 1 - err / meta) if meta else None
    desv = (si - meta) / meta if meta else None
    return {"meta": round(meta, 1), "si": round(si, 1), "err": round(err, 1),
            "exact": round(exact, 4) if exact is not None else None,
            "desv": round(desv, 4) if desv is not None else None,
            "n_sku": len({f["sku"] for f in filas})}


def group_by(filas, key):
    out = defaultdict(list)
    for f in filas:
        out[f[key]].append(f)
    return out


def main():
    if len(sys.argv) != 2:
        print(__doc__)
        sys.exit(1)

    filas = leer_filas(sys.argv[1])  # MENSUAL + SEMANAL combinados = total real del mes

    meses = sorted({f["mes"] for f in filas})
    ultimo_mes = meses[-1]
    anio_actual = ultimo_mes[:4]
    meses_anio = [m for m in meses if m.startswith(anio_actual)]

    # A) Evolucion mensual del anio actual (total compania)
    evolucion = [{"mes": mes, **agg([f for f in filas if f["mes"] == mes])} for mes in meses_anio]

    # B) Cadena ranking del ultimo mes
    filas_ultimo_mes = [f for f in filas if f["mes"] == ultimo_mes]
    por_cadena = group_by(filas_ultimo_mes, "cadena")
    cadena_ranking = [{"cadena": c, **agg(fs)} for c, fs in por_cadena.items()]
    cadena_ranking.sort(key=lambda r: (r["exact"] if r["exact"] is not None else 1))

    # C) Matriz CPFR x mes (anio actual)
    cpfrs = sorted({f["cpfr"] for f in filas if f["cpfr"]})
    matriz_cpfr = {}
    for cpfr in cpfrs:
        fila = {}
        for mes in meses_anio:
            fs = [f for f in filas if f["mes"] == mes and f["cpfr"] == cpfr]
            fila[mes] = agg(fs) if fs else None
        matriz_cpfr[cpfr] = fila

    # D) Detalle por categoria: acumulado del anio + mes actual, filtrable por CPFR en el HTML
    filas_anio = [f for f in filas if f["mes"].startswith(anio_actual)]

    def detalle_categoria(filas_base):
        por_cat = group_by(filas_base, "categoria")
        return sorted(
            [{"categoria": c, **agg(fs)} for c, fs in por_cat.items() if c],
            key=lambda r: (r["exact"] if r["exact"] is not None else 1)
        )

    detalle = {"TODAS": {"acumulado": detalle_categoria(filas_anio), "mes_actual": detalle_categoria(filas_ultimo_mes)}}
    for cpfr in cpfrs:
        detalle[cpfr] = {
            "acumulado": detalle_categoria([f for f in filas_anio if f["cpfr"] == cpfr]),
            "mes_actual": detalle_categoria([f for f in filas_ultimo_mes if f["cpfr"] == cpfr]),
        }

    # E) Top SKU mas desviados por cadena (mayor |desviacion|), acumulado y mes actual
    TOP_N_SKU = 15
    cadenas = sorted({f["cadena"] for f in filas if f["cadena"]})

    def top_sku_cadena(filas_base, cadena):
        por_sku = group_by([f for f in filas_base if f["cadena"] == cadena], "sku")
        filas_sku = []
        for sku, fs in por_sku.items():
            if not sku:
                continue
            a = agg(fs)
            filas_sku.append({
                "sku": sku, "nombre": fs[0]["nombre"], "categoria": fs[0]["categoria"], **a,
            })
        filas_sku.sort(key=lambda r: abs(r["desv"]) if r["desv"] is not None else -1, reverse=True)
        return filas_sku[:TOP_N_SKU]

    top_sku_por_cadena = {
        "acumulado": {c: top_sku_cadena(filas_anio, c) for c in cadenas},
        "mes_actual": {c: top_sku_cadena(filas_ultimo_mes, c) for c in cadenas},
    }

    data = {
        "generado_desde_mes": ultimo_mes,
        "meta_objetivo": META_OBJETIVO,
        "evolucion": evolucion,
        "cadena_ranking": cadena_ranking,
        "matriz_cpfr": matriz_cpfr,
        "meses_matriz": meses_anio,
        "cpfrs": cpfrs,
        "detalle": detalle,
        "cadenas": cadenas,
        "top_sku_por_cadena": top_sku_por_cadena,
        "total_acumulado": agg(filas_anio),
        "total_mes_actual": agg(filas_ultimo_mes),
    }

    (HERE / "exactitud_data.json").write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")

    template = (HERE / "dashboard_exactitud_template.html").read_text(encoding="utf-8")
    html = template.replace("__DATA_JSON__", json.dumps(data, ensure_ascii=False))
    OUT_HTML.write_text(html, encoding="utf-8")
    print(f"Dashboard generado: {OUT_HTML} ({OUT_HTML.stat().st_size/1024:.0f} KB)")
    print(f"Ultimo mes (en curso): {ultimo_mes} | Meses en evolucion: {len(meses_anio)} | CPFR: {len(cpfrs)}")


if __name__ == "__main__":
    main()
