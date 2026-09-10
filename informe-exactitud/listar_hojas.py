"""
Paso 1: inspecciona el Excel maestro de Exactitud y muestra que hojas tiene
(incluidas las ocultas), para saber como identificar la semana de cada una
antes de escribir el exportador automatico.

Uso:
    python listar_hojas.py "Informe Exactitud W35-2026.xlsx"
"""
import sys
import re
from openpyxl import load_workbook


def buscar_semana_en_celdas(ws, max_filas=40):
    """Busca celdas con el texto 'Semana' y devuelve el valor de la celda siguiente."""
    encontrados = []
    for row in ws.iter_rows(min_row=1, max_row=max_filas):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip() == "Semana":
                vecina = ws.cell(row=cell.row, column=cell.column + 1).value
                encontrados.append(vecina)
    return encontrados


def main(path):
    wb = load_workbook(path, data_only=True)
    print(f"Archivo: {path}")
    print(f"Total hojas: {len(wb.sheetnames)}\n")
    print(f"{'Nombre hoja':30} {'Estado':10} {'Semana(s) detectada(s) en celdas'}")
    print("-" * 90)
    for name in wb.sheetnames:
        ws = wb[name]
        estado = ws.sheet_state  # 'visible', 'hidden', 'veryHidden'
        try:
            semanas = buscar_semana_en_celdas(ws)
        except Exception as e:
            semanas = [f"ERROR: {e}"]
        print(f"{name[:30]:30} {estado:10} {semanas}")

    print("\nListo. Copia y pega toda esta salida en el chat.")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print(__doc__)
        sys.exit(1)
    main(sys.argv[1])
